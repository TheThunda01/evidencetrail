"""Read bounded CSV or XLSX uploads as literal extraction data."""

import csv
from io import BytesIO, StringIO
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml import ElementTree as safe_xml
from defusedxml.common import DefusedXmlException
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .audit import MAX_COLUMNS, MAX_ROWS, normalize_headers

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 25 * 1024 * 1024


def _check_size(content: bytes) -> None:
    if not isinstance(content, bytes):
        raise ValueError("Provide the file contents as bytes.")
    if not content:
        raise ValueError("The file is empty. Add a header row and try again.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("The file is too large. The upload limit is 5 MiB.")


def _open_xlsx(content: bytes):
    _check_size(content)
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ValueError("The expanded workbook is too large. The limit is 25 MiB.")
            # Preflight every XML member before openpyxl reads the archive. This
            # remains protective even if its optional safe parser is disabled.
            # Clear completed elements to keep preflight memory bounded.
            for member in archive.infolist():
                with archive.open(member) as source:
                    prefix = source.read(256)
                    # Relationships can point to XML with a nonstandard file
                    # extension. Recognize XML signatures too, including BOMs
                    # and the zero bytes in UTF-16/UTF-32 encodings.
                    signature = prefix.lstrip(b" \t\r\n\x00\xef\xbb\xbf\xff\xfe")
                    while prefix and not signature:
                        prefix = source.read(4096)
                        signature = prefix.lstrip(b" \t\r\n\x00\xef\xbb\xbf\xff\xfe")
                    looks_like_xml = signature.startswith(b"<")
                    if not member.filename.lower().endswith((".xml", ".rels")) and not looks_like_xml:
                        continue
                    source.seek(0)
                    for _, element in safe_xml.iterparse(
                        source, events=("end",), forbid_dtd=True,
                        forbid_entities=True, forbid_external=True,
                    ):
                        element.clear()
    except DefusedXmlException as exc:
        raise ValueError("This workbook contains unsupported XML declarations or entities. Save a new XLSX copy and try again.") from exc
    except ParseError as exc:
        raise ValueError("This file could not be read as an XLSX workbook. Save it as .xlsx and try again.") from exc
    except (BadZipFile, InvalidFileException, KeyError, OSError, EOFError, RuntimeError, NotImplementedError) as exc:
        raise ValueError("This file could not be read as an XLSX workbook. Save it as .xlsx and try again.") from exc
    try:
        return load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        # Optional XML parsers raise different exception classes. Contain those
        # errors at the upload boundary without displaying parser internals.
        raise ValueError("This file could not be read as an XLSX workbook. Save it as .xlsx and try again.") from exc


def list_sheets(content: bytes) -> list[str]:
    """List worksheet names in a size-checked XLSX upload without changing it."""
    workbook = _open_xlsx(content)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _build_frame(headers: list[object], rows: list[list[object]]) -> pd.DataFrame:
    if not headers:
        raise ValueError("The table has no header row.")
    if len(headers) > MAX_COLUMNS:
        raise ValueError(f"The table has too many columns. The limit is {MAX_COLUMNS}.")
    if len(rows) > MAX_ROWS:
        raise ValueError(f"The table has too many records. The limit is {MAX_ROWS:,}.")
    return pd.DataFrame(rows, columns=normalize_headers(headers))


def load_table(filename: str, content: bytes, sheet_name: str | None = None) -> pd.DataFrame:
    """Load one table, preserving literal text such as NA and null.

    CSV must use comma separators and UTF-8 (an optional BOM is accepted). XLSX
    uses the requested sheet, or the first sheet by default. XLSX formulas in the
    selected sheet are rejected: export values to a separate file before audit.
    No formulas, macros, URLs, or external workbook links are executed.
    """
    _check_size(content)
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("The CSV must use UTF-8 encoding. Save a UTF-8 CSV and try again.") from exc
        if "\x00" in text:
            raise ValueError("The CSV contains null bytes and cannot be read as a text table.")
        try:
            reader = csv.reader(StringIO(text, newline=""), strict=True)
            headers = next(reader, [])
            if len(headers) > MAX_COLUMNS:
                raise ValueError(f"The table has too many columns. The limit is {MAX_COLUMNS}.")
            normalize_headers(headers)
            rows = []
            for row_number, row in enumerate(reader, start=2):
                if not row:
                    row = [""] * len(headers)
                if len(row) != len(headers):
                    raise ValueError(f"CSV record {row_number} has {len(row)} fields; the header has {len(headers)}.")
                rows.append(row)
                if len(rows) > MAX_ROWS:
                    raise ValueError(f"The table has too many records. The limit is {MAX_ROWS:,}.")
        except csv.Error as exc:
            raise ValueError("The CSV could not be read. Check its quoting, separators, and cell lengths.") from exc
        return _build_frame(headers, rows)
    if suffix != ".xlsx":
        raise ValueError("Choose a .csv or .xlsx file. Other file types are not supported.")

    workbook = _open_xlsx(content)
    try:
        if not workbook.sheetnames:
            raise ValueError("The workbook has no worksheets to read.")
        selected = sheet_name if sheet_name is not None else workbook.sheetnames[0]
        if selected not in workbook.sheetnames:
            raise ValueError("The selected worksheet is not in this workbook.")
        sheet = workbook[selected]
        if sheet.max_column is not None and sheet.max_column > MAX_COLUMNS:
            raise ValueError(f"The worksheet has too many columns. The limit is {MAX_COLUMNS}.")
        if sheet.max_row is not None and sheet.max_row > MAX_ROWS + 1:
            raise ValueError(f"The worksheet has too many records. The limit is {MAX_ROWS:,}.")
        table_rows = []
        for row_number, cells in enumerate(sheet.iter_rows(), start=1):
            if row_number > MAX_ROWS + 1:
                raise ValueError(f"The worksheet has too many records. The limit is {MAX_ROWS:,}.")
            if len(cells) > MAX_COLUMNS:
                raise ValueError(f"The worksheet has too many columns. The limit is {MAX_COLUMNS}.")
            for cell in cells:
                if cell.data_type == "f":
                    raise ValueError(f"The selected worksheet contains a formula at {cell.coordinate}. Export values to a separate file before auditing.")
            table_rows.append(["" if cell.value is None else cell.value for cell in cells])
        if not table_rows or all(value == "" for value in table_rows[0]):
            raise ValueError("The selected worksheet has no header row.")
        return _build_frame(table_rows[0], table_rows[1:])
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("The selected worksheet could not be read. Save a new XLSX copy and try again.") from exc
    finally:
        workbook.close()
