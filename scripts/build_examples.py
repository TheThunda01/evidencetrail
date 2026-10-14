"""Rebuild the entirely synthetic CSV and Excel examples."""

import csv
from copy import deepcopy
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
COLUMNS = ["record_id", "report_id", "field", "value", "evidence_status", "source_locator", "review_status", "reviewer", "notes"]
FIELDS = ["Responsible organization", "Data sources", "Information exchange", "Public health action"]
VALUES = ["Fictional district health team", "Synthetic human and animal case reports", "Weekly summary shared with a fictional partner", "Synthetic example of a joint review meeting"]


def write_csv(path, rows):
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def main():
    destination = ROOT / "examples"
    destination.mkdir(exist_ok=True)
    clean = []
    for i in range(20):
        clean.append(dict(zip(COLUMNS, [
            f"R{i + 1:03d}", f"SIM-{i // 4 + 1:03d}", FIELDS[i % 4], VALUES[i % 4],
            "Reported", f"Synthetic source, section {i % 4 + 1}", "Pending", "Demo reviewer",
            "Fictional record for software testing. Not a real study.",
        ])))
    for index, status in [(9, "Not reported"), (13, "Unclear"), (17, "Not applicable")]:
        clean[index].update(value="", evidence_status=status, source_locator="")
    clean[11].update(value="The fictional example explicitly states no shared review took place.", evidence_status="Explicitly absent")
    demo = deepcopy(clean)
    demo[3]["record_id"] = "R003"
    demo[4]["source_locator"] = ""
    demo[5]["evidence_status"] = "Not stated"
    demo[6]["value"] = ""
    demo[7]["review_status"] = "Disagreement"
    write_csv(destination / "demo_records.csv", demo)
    write_csv(destination / "clean_records.csv", clean)
    write_csv(destination / "reviewer_a.csv", clean)
    second_reviewer = deepcopy(clean)
    second_reviewer[0]["value"] = "Fictional regional health team"
    second_reviewer[2].update(value="", evidence_status="Unclear", source_locator="")
    second_reviewer[4]["source_locator"] = "Synthetic source, table 2"
    second_reviewer.pop()
    extra = deepcopy(clean[0])
    extra.update(record_id="R021", report_id="SIM-006")
    second_reviewer.append(extra)
    write_csv(destination / "reviewer_b.csv", second_reviewer)
    write_csv(destination / "extraction_template.csv", [])
    expected = [
        {"code": "DUPLICATE_RECORD_ID", "row_numbers": [4, 5]},
        {"code": "MISSING_SOURCE_LOCATOR", "row_numbers": [6]},
        {"code": "INVALID_EVIDENCE_STATUS", "row_numbers": [7]},
        {"code": "MISSING_REPORTED_VALUE", "row_numbers": [8]},
        {"code": "REVIEWER_DISAGREEMENT", "row_numbers": [9]},
    ]
    (destination / "expected_findings.json").write_text(json.dumps(expected, indent=2) + "\n", encoding="utf-8")

    book = Workbook()
    book.remove(book.active)
    for name, rows in [("Demo", demo), ("Clean", clean)]:
        sheet = book.create_sheet(name)
        sheet.append(COLUMNS)
        for row in rows:
            sheet.append([row[column] for column in COLUMNS])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 28
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="176B57")
            cell.font = Font(color="FFFFFF", bold=True)
        widths = [15, 16, 30, 58, 21, 39, 21, 22, 52]
        for column, width in zip("ABCDEFGHI", widths):
            sheet.column_dimensions[column].width = width
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 46
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column, choices in [("E", "Reported,Not reported,Unclear,Explicitly absent,Not applicable"), ("G", "Pending,Verified,Disagreement")]:
            validation = DataValidation(type="list", formula1=f'"{choices}"')
            validation.showErrorMessage = True
            validation.error = "Choose one of the documented status values."
            sheet.add_data_validation(validation)
            validation.add(f"{column}2:{column}10001")
    book.save(destination / "demo_records.xlsx")


if __name__ == "__main__":
    main()
